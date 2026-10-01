"""Producing a document from a validated selection.

This is where the module's promise is either kept or broken. Everything the
platform has built — the ranking, the frozen shortlist, the human decision, the
template library — exists so that this function can put a real consultant's
real experience into a funder's real form, and so that every field in the
result traces back to a row somebody can point at.

Two rules govern the context built here, and both are refusals.

**Nothing is invented.** Every value comes from stored data: the CV's own
extracted timeline, the criteria read from it, the tender it is being submitted
for, the shortlist entry that retained it. No field is inferred, defaulted to
something plausible, or filled from a model. A missing employer stays missing,
because "we could not read it" and "there wasn't one" are different facts and
only the first is true.

**Nothing is generated for someone nobody chose.** The input is a *retained*
entry of a *validated* shortlist. A profile left pending was not selected; it
was not looked at. Producing a CV for them would put a real person's name on a
real submission on the strength of a rank.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime
from typing import Any

from app.core.identity import utc_now
from app.core.logging import get_logger

logger = get_logger(__name__)

__all__ = [
    "GenerationRefused",
    "RenderedDocument",
    "build_cv_context",
    "generate_cv",
]

#: Stamped on every version that has not been approved. The parcours asks for
#: it explicitly, and the reason is not ceremony: a draft that looks final is a
#: draft that gets sent.
WATERMARK = "GÉNÉRÉ — EN ATTENTE DE VALIDATION"

#: How many roles reach the document. A funder's form has a box, not a chapter,
#: and the extractor has been measured returning up to twenty entries for one
#: CV. The most recent are the ones a reviewer reads.
_MAX_EXPERIENCES = 8
_MAX_FORMATIONS = 6


class GenerationRefused(Exception):
    """A document that must not be produced, with the reason a human needs."""


@dataclass(slots=True)
class RenderedDocument:
    filename: str
    content: bytes
    #: Variables the template asked for and the data could not fill. Reported
    #: rather than hidden: the QA pass in the next step turns these into a
    #: checklist, and until it exists a human is owed the same list.
    empty_fields: list[str]
    context_keys: list[str]
    #: What the reformulation pass did, or why it did nothing. Carried out of
    #: here so a produced document can state whether a model touched it —
    #: which a reviewer is entitled to know before signing.
    adaptation: dict[str, Any] = field(default_factory=dict)
    #: The automatic review of the file itself: what was found, what blocked,
    #: and which paragraphs had to be put back. A document that passed and a
    #: document nobody checked are different things, and only one of them
    #: should reach a reviewer unannounced.
    qa: dict[str, Any] = field(default_factory=dict)


def build_cv_context(
    *,
    cv: Any,
    entry: Any = None,
    shortlist: Any = None,
    tender: Any = None,
    approved: bool = False,
    now: datetime | None = None,
    adapter: Any = None,
) -> tuple[dict[str, Any], dict[str, Any], list[dict[str, str]]]:
    """Assemble what a CV template is filled from.

    Takes objects rather than ids and opens no session, so the rules are
    testable without a database — which is what keeps them tested.

    ``approved`` empties the watermark. It is a parameter rather than a lookup
    so that the only way to produce an unmarked document is for a caller to
    state, in one word, that it was approved.

    ``adapter(experiences, requirements, limit) -> (experiences, report)``
    runs the roles through the reformulation pass. It changes the *wording* of
    missions and their order, never their content: every reformulated passage
    is checked back against its source and discarded if it names a technology,
    an acronym or a year the source did not, or if it has drifted off subject.

    It is injected rather than imported because the check needs the encoder,
    and the encoder must not be loaded in the API. Omitted, or failing for any
    reason, the deterministic text is what reaches the document.

    Returns the context, a report of what the adaptation did, and the roles
    as they were *before* adaptation. The third is what the QA pass puts back
    when a reformulated paragraph fails: a correction needs somewhere correct
    to fall back to.
    """
    moment = now or utc_now()
    criteria = dict(getattr(cv, "criteria", None) or {})
    structure = dict(getattr(cv, "structure", None) or {})

    raw_experiences = [_experience(item) for item in (structure.get("experiences") or [])]
    adaptation: dict[str, Any] = {"status": "off", "llm_used": False}
    experiences = raw_experiences[:_MAX_EXPERIENCES]
    if adapter is not None and raw_experiences:
        try:
            # The requirements come from the *frozen* shortlist, not from a
            # fresh read of the tender: the wording a role is aligned to must
            # be the wording the selection was judged against.
            adapted, adaptation = adapter(
                raw_experiences, _requirement_texts(shortlist), _MAX_EXPERIENCES
            )
            experiences = adapted or experiences
        except Exception as exc:
            # A worker that is down, a timed-out broker, a malformed answer.
            # The deterministic document is produced either way: the pass is an
            # improvement, never a dependency.
            logger.warning("generation.adaptation_failed", error=str(exc)[:200])
            adaptation = {"status": "failed", "llm_used": False}
    formations = [
        _formation(item) for item in (structure.get("formations") or [])[:_MAX_FORMATIONS]
    ]

    # The label the validator saw, not a fresh lookup. A CV re-imported under a
    # different display name must not silently change the name on a document
    # produced from a selection someone already approved.
    name = (
        getattr(entry, "label", None)
        or getattr(cv, "display_name", None)
        or getattr(cv, "headline", None)
        or getattr(cv, "original_filename", "")
    )

    context: dict[str, Any] = {
        "nom": name,
        "poste": getattr(cv, "headline", None) or "",
        "fichier": getattr(cv, "original_filename", "") or "",
        "competences": _strings(criteria.get("technologies")),
        "langues": _strings(criteria.get("languages")),
        "certifications": _strings(criteria.get("certifications")),
        "niveau_etudes": criteria.get("education_label") or "",
        "formations": formations,
        "experiences": experiences,
        "mission_titre": getattr(tender, "title", None)
        or getattr(shortlist, "tender_title", None)
        or "",
        "mission_acheteur": getattr(tender, "buyer", None) or "",
        "mission_reference": getattr(tender, "reference", None) or "",
        "mission_pays": getattr(tender, "country", None) or "",
        "mission_echeance": _date(getattr(tender, "deadline", None)),
        "rang": getattr(entry, "rank", "") or "",
        "score": f"{getattr(entry, 'score', 0.0):.2f}" if entry is not None else "",
        # The passages that justified the ranking, as they were scored. Copied
        # from the frozen entry rather than recomputed: the document must show
        # the evidence the decision was taken on.
        "extraits": _evidence(entry),
        "genere_le": moment.strftime("%d/%m/%Y"),
        "filigrane": "" if approved else WATERMARK,
    }
    return context, adaptation, raw_experiences


def generate_cv(
    *,
    template_bytes: bytes,
    cv: Any,
    entry: Any = None,
    shortlist: Any = None,
    tender: Any = None,
    approved: bool = False,
    now: datetime | None = None,
    adapter: Any = None,
    kind: str = "cv",
) -> RenderedDocument:
    """Fill one CV template.

    Deterministic unless an ``adapter`` is given. Even then the model only
    rewords and reorders what the CV already says: the layout, the dates, the
    identity and every fact come from stored rows, and a reformulation that
    adds a technology, an acronym or a year — or that drifts off subject — is
    thrown away before it reaches the file.
    """
    from app.services.templates import TemplateError, render_template

    if entry is not None and getattr(entry, "decision", None) not in (None, "retained"):
        raise GenerationRefused(
            "Ce profil n'a pas été retenu. Un document ne peut être produit que "
            "pour une décision explicite."
        )

    context, adaptation, source = build_cv_context(
        cv=cv,
        entry=entry,
        shortlist=shortlist,
        tender=tender,
        approved=approved,
        now=now,
        adapter=adapter,
    )

    def _render(ctx: dict[str, Any]) -> bytes:
        try:
            return render_template(template_bytes, ctx)
        except TemplateError as exc:
            raise GenerationRefused(str(exc)) from exc

    from app.services.document_qa import repair_context, review

    content = _render(context)
    report = review(_document_text(content), context=context, kind=kind)
    repaired: list[str] = []

    if not report.passed:
        # Targeted correction: the failing paragraphs go back to what the CV
        # actually says, and the document is rendered once more. At most two
        # passes, because the fallback text passes these checks by
        # construction — there is no third outcome.
        context, repaired = repair_context(context, report, source)
        if repaired:
            content = _render(context)
            report = review(_document_text(content), context=context, kind=kind)

    empty = sorted(key for key, value in context.items() if key != "filigrane" and not value)
    logger.info(
        "generation.cv_produced",
        bytes=len(content),
        experiences=len(context["experiences"]),
        formations=len(context["formations"]),
        empty_fields=len(empty),
        qa_passed=report.passed,
        qa_blocking=len(report.blocking),
        repaired=len(repaired),
        **{f"adaptation_{k}": v for k, v in adaptation.items() if k != "invented_terms"},
    )
    return RenderedDocument(
        filename=_filename(context["nom"], context["mission_reference"]),
        content=content,
        empty_fields=empty,
        context_keys=sorted(context),
        adaptation=adaptation,
        qa={**report.to_dict(), "repaired_sections": repaired},
    )


# ---------------------------------------------------------------------------
def _document_text(content: bytes) -> str:
    """The document's own words.

    The QA pass reads the artefact rather than the context on purpose: a
    template that silently dropped a section produces a file the context
    cannot describe, and that is the failure most worth catching.
    """
    import io

    try:
        from docx import Document

        document = Document(io.BytesIO(content))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            parts.extend(cell.text for row in table.rows for cell in row.cells)
        return "\n".join(parts)
    except Exception:
        # Unreadable output is itself a finding, and `review` will raise it as
        # one: an empty string carries no name and no sections.
        return ""


def _requirement_texts(shortlist: Any) -> list[str]:
    """The requirement passages the selection was frozen with."""
    requirements = getattr(shortlist, "requirements", None) or []
    out: list[str] = []
    for item in requirements:
        text = item.get("text") if isinstance(item, dict) else None
        if text:
            out.append(str(text)[:300])
    return out


def _experience(item: dict[str, Any]) -> dict[str, str]:
    """One role, with every field a string the template can print.

    ``None`` is turned into ``""`` here rather than in the template, so a Word
    file written by someone who is not a developer never has to know that a
    missing employer is a different thing from an empty one.
    """
    return {
        "debut": str(item.get("debut") or ""),
        "fin": str(item.get("fin") or ""),
        "poste": str(item.get("poste") or ""),
        "employeur": str(item.get("employeur") or ""),
        "lieu": str(item.get("lieu") or ""),
        "missions": str(item.get("missions") or ""),
    }


def _formation(item: dict[str, Any]) -> dict[str, str]:
    return {
        "annee": str(item.get("annee") or ""),
        "diplome": str(item.get("diplome") or ""),
        "etablissement": str(item.get("etablissement") or ""),
    }


def _strings(value: Any) -> list[str]:
    if not value:
        return []
    return [str(item) for item in value if str(item).strip()]


def _evidence(entry: Any) -> list[str]:
    passages = getattr(entry, "evidence", None) or []
    out: list[str] = []
    for item in passages[:4]:
        text = item.get("passage") if isinstance(item, dict) else None
        if text:
            out.append(str(text).strip()[:400])
    return out


def _date(value: Any) -> str:
    if isinstance(value, datetime):
        return value.strftime("%d/%m/%Y")
    return str(value) if value else ""


def _filename(name: str, reference: str) -> str:
    """A name a bid manager can find in a folder of forty.

    Built from the person and the tender, not from a uuid: the file is going
    into someone's dossier, and ``cv_9f3a....docx`` is a file nobody opens.
    """
    parts = [_slug(name) or "cv", _slug(reference)]
    return "_".join(part for part in parts if part)[:120] + ".docx"


def _slug(value: str) -> str:
    import unicodedata

    normalised = unicodedata.normalize("NFKD", str(value or ""))
    ascii_only = normalised.encode("ascii", "ignore").decode("ascii")
    cleaned = "".join(c if c.isalnum() else "_" for c in ascii_only)
    return "_".join(part for part in cleaned.split("_") if part)[:60]
