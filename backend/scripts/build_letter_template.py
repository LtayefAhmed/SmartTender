"""Build the default covering-letter template.

The simplest of the three layouts, and the one where the template carries the
least: the body is three paragraphs the platform composed, so the file is a
letterhead, an object line and a loop.

That asymmetry is worth noticing. For a CV the template does most of the work
— it decides which of nineteen variables appear and in what order. Here it
decides almost nothing, because what the letter *says* is the whole document.
Which is also why this is the one kind where the model can do real damage and
the guard refuses the draft whole rather than mending a paragraph.

Run from ``backend/``::

    python scripts/build_letter_template.py
"""

from __future__ import annotations

import argparse
from pathlib import Path

DEFAULT_OUT = Path("config/templates/lettre_accompagnement.docx")


def build(destination: Path) -> Path:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt, RGBColor

    document = Document()

    for section in document.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        for margin in ("top_margin", "bottom_margin"):
            setattr(section, margin, Cm(2.4))
        for margin in ("left_margin", "right_margin"):
            setattr(section, margin, Cm(2.5))

    document.styles["Normal"].font.name = "Calibri"
    document.styles["Normal"].font.size = Pt(10.5)

    banner = document.add_paragraph()
    banner.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = banner.add_run("{{ filigrane }}")
    run.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0xB0, 0x30, 0x30)

    # Addressee, right-aligned as a letter is.
    addressee = document.add_paragraph()
    addressee.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    addressee.add_run("{{ mission_acheteur }}")

    date = document.add_paragraph()
    date.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    date.add_run("Le {{ genere_le }}")

    document.add_paragraph()

    subject = document.add_paragraph()
    subject.add_run("Objet : ").bold = True
    subject.add_run(
        "candidature à la consultation « {{ mission_titre }} »"
        "{% if mission_reference %} — référence {{ mission_reference }}{% endif %}"
    )

    document.add_paragraph()
    document.add_paragraph("Madame, Monsieur,")
    document.add_paragraph()

    # The body. One paragraph per entry, with a blank line between them —
    # `{%p %}` repeats the whole block, so the spacing travels with the text.
    document.add_paragraph("{%p for p in paragraphes %}")
    document.add_paragraph("{{ p }}")
    document.add_paragraph()
    document.add_paragraph("{%p endfor %}")

    document.add_paragraph(
        "Nous vous prions d'agréer, Madame, Monsieur, l'expression de nos "
        "salutations distinguées."
    )

    document.add_paragraph()
    document.add_paragraph()
    signature = document.add_paragraph()
    signature.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    # Left blank on purpose: the platform does not sign. A human does, and the
    # parcours is explicit that responsibility stays with a person.
    signature.add_run("Signature et qualité du signataire").italic = True

    document.add_paragraph()
    footer = document.add_paragraph()
    run = footer.add_run(
        "Pièces jointes : curriculum vitæ des intervenants, matrice de conformité."
    )
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    destination.parent.mkdir(parents=True, exist_ok=True)
    staging = destination.with_suffix(".docx.part")
    document.save(staging)
    staging.replace(destination)
    return destination


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", type=Path, default=DEFAULT_OUT)
    args = parser.parse_args()

    path = build(args.out)

    from app.services.templates import check_variables, discover_variables

    found = discover_variables(path.read_bytes())
    unknown, unavailable = check_variables(found, kind="lettre")

    print(f"écrit      {path} ({path.stat().st_size:,} octets)")
    print(f"variables  {', '.join(sorted(found))}")
    if unavailable:
        print(f"en attente {', '.join(unavailable)}")
    if unknown:
        print(f"INCONNUES  {', '.join(unknown)}")
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
