"""Build the default expert-sheet template.

A *fiche expert* is not a short CV. It is the page that goes into the technical
offer to present one consultant to the client: who they are, what they bring to
*this* mission, and three experiences that prove it. A full CV is the annex; the
sheet is the argument.

It runs on exactly the same render context as the CV — same nineteen variables,
same data, same code path. Only the template differs. That is the whole claim
of "templates are data" made concrete: a new document type here costs a layout,
not a release.

Two deliberate omissions, and they are the interesting part:

**No rank, no matching score.** The CV template prints them as provenance, for
an internal reader who may want to know where the document came from. This
sheet is handed to the buyer, and "score de rapprochement 0,59" is an internal
number that would raise a question nobody wants asked in a bid defence. The
variables exist; this template simply does not use them.

**Three experiences, not eight.** The render context supplies up to eight,
already ordered by relevance. The slice happens here because it is a
presentation choice, not a data one — the CV annex still shows all of them.

Run from ``backend/``::

    python scripts/build_expert_template.py
"""

from __future__ import annotations

import argparse
from pathlib import Path

DEFAULT_OUT = Path("config/templates/fiche_expert.docx")


def build(destination: Path) -> Path:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt, RGBColor

    document = Document()

    for section in document.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        for margin in ("top_margin", "bottom_margin"):
            setattr(section, margin, Cm(1.8))
        for margin in ("left_margin", "right_margin"):
            setattr(section, margin, Cm(2.0))

    document.styles["Normal"].font.name = "Calibri"
    document.styles["Normal"].font.size = Pt(10)

    banner = document.add_paragraph()
    banner.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = banner.add_run("{{ filigrane }}")
    run.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0xB0, 0x30, 0x30)

    heading = document.add_paragraph()
    run = heading.add_run("FICHE EXPERT")
    run.bold = True
    run.font.size = Pt(15)

    name = document.add_paragraph()
    run = name.add_run("{{ nom }}")
    run.bold = True
    run.font.size = Pt(13)

    # Only when it adds something. On an anonymised CV the platform has no
    # name to print, so `nom` falls back to the job title — and the sheet
    # showed "AMC COMPUTER SPECIALIST" twice, once as the person and once as
    # the role. The condition costs nothing and removes a line that reads as a
    # bug to anyone who opens the file.
    role = document.add_paragraph()
    run = role.add_run("{% if poste and poste != nom %}{{ poste }}{% endif %}")
    run.italic = True
    run.font.color.rgb = RGBColor(0x50, 0x50, 0x50)

    # The mission on the face of the sheet: the same expert is presented
    # differently for a TMA and for an audit, and a page that does not name
    # the tender gets filed against the wrong one.
    mission = document.add_paragraph()
    run = mission.add_run(
        "Proposé pour : {{ mission_titre }}"
        "{% if mission_acheteur %} — {{ mission_acheteur }}{% endif %}"
        "{% if mission_reference %} ({{ mission_reference }}){% endif %}"
    )
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x50, 0x50, 0x50)

    _title(document, "Compétences clés")
    document.add_paragraph("{{ competences | join(' · ') }}")

    facts = document.add_table(rows=0, cols=2)
    facts.style = "Table Grid"
    for label, placeholder in (
        ("Niveau d'études", "{{ niveau_etudes }}"),
        ("Langues", "{{ langues | join(', ') }}"),
        ("Certifications", "{{ certifications | join(', ') }}"),
    ):
        cells = facts.add_row().cells
        cells[0].paragraphs[0].add_run(label).bold = True
        cells[1].text = placeholder

    # Three, sliced in the template. The context supplies eight, already
    # ordered by relevance to this tender; how many reach the page is a
    # presentation decision and belongs here, not in the code.
    _title(document, "Expériences les plus proches de la mission")
    document.add_paragraph("{%p for e in experiences[:3] %}")
    header = document.add_paragraph()
    header.add_run(
        "{{ e.debut }} – {{ e.fin }} · {{ e.poste }}"
        "{% if e.employeur %}, {{ e.employeur }}{% endif %}"
    ).bold = True
    document.add_paragraph("{{ e.missions }}")
    document.add_paragraph("{%p endfor %}")

    _title(document, "Formation")
    document.add_paragraph("{%p for f in formations %}")
    document.add_paragraph("{{ f.annee }} — {{ f.diplome }}, {{ f.etablissement }}")
    document.add_paragraph("{%p endfor %}")

    # No rank, no score: this page is read by the buyer.
    document.add_paragraph()
    footer = document.add_paragraph()
    run = footer.add_run("Généré le {{ genere_le }} · CV détaillé en annexe")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    destination.parent.mkdir(parents=True, exist_ok=True)
    staging = destination.with_suffix(".docx.part")
    document.save(staging)
    staging.replace(destination)
    return destination


def _title(document, text: str) -> None:
    from docx.shared import Pt

    paragraph = document.add_paragraph()
    run = paragraph.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(10.5)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", type=Path, default=DEFAULT_OUT)
    args = parser.parse_args()

    path = build(args.out)

    from app.services.templates import check_variables, discover_variables

    found = discover_variables(path.read_bytes())
    unknown, unavailable = check_variables(found, kind="fiche_expert")

    print(f"écrit      {path} ({path.stat().st_size:,} octets)")
    print(f"variables  {', '.join(sorted(found))}")
    if unavailable:
        print(f"en attente {', '.join(unavailable)}")
    if unknown:
        print(f"INCONNUES  {', '.join(unknown)}")
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
