"""Build the default CV template.

The template is generated from this file rather than committed as a ``.docx``,
and the reason is reviewability. A Word document in git is an opaque binary: a
reviewer sees "template.docx changed, 40 KB", and nobody can tell whether a
placeholder was renamed, a section dropped, or a macro added. Here the layout
*is* the diff.

It also means the template can be rebuilt from scratch on any machine, which
matters for the handover: a corrupted or lost file is one command away, not a
request to whoever last had it.

The layout follows the structure common to funder CV forms — BEI, BAD, Bureau
d'Études all ask for the same skeleton in a different order: who the person is,
what the mission is, formal education dated, languages, then the career
history with the missions spelled out. When a real imposed form arrives it is
uploaded and supersedes this one; nothing here is load-bearing.

Run from ``backend/``::

    python scripts/build_default_template.py
    python scripts/build_default_template.py --out /tmp/cv.docx
"""

from __future__ import annotations

import argparse
from pathlib import Path

DEFAULT_OUT = Path("config/templates/cv_inetum.docx")


def build(destination: Path) -> Path:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    document = Document()

    # A4 with sane margins: a funder's form is printed and filed, and the
    # default US Letter would reflow every page in a European office.
    from docx.shared import Cm

    for section in document.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        for margin in ("top_margin", "bottom_margin"):
            setattr(section, margin, Cm(2.0))
        for margin in ("left_margin", "right_margin"):
            setattr(section, margin, Cm(2.2))

    base = document.styles["Normal"]
    base.font.name = "Calibri"
    base.font.size = Pt(10.5)

    # --- state banner ------------------------------------------------------
    # Rendered empty once the document is approved. A version awaiting
    # validation has to say so on its own face: a draft that looks final is how
    # a draft gets sent to a client.
    banner = document.add_paragraph()
    banner.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = banner.add_run("{{ filigrane }}")
    run.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0xB0, 0x30, 0x30)

    # --- identity ----------------------------------------------------------
    heading = document.add_paragraph()
    run = heading.add_run("CURRICULUM VITAE")
    run.bold = True
    run.font.size = Pt(16)

    identity = document.add_table(rows=0, cols=2)
    identity.style = "Table Grid"
    for label, placeholder in (
        ("Nom", "{{ nom }}"),
        ("Poste proposé", "{{ poste }}"),
        ("Niveau d'études", "{{ niveau_etudes }}"),
    ):
        cells = identity.add_row().cells
        cells[0].paragraphs[0].add_run(label).bold = True
        cells[1].text = placeholder

    # --- the mission this CV is submitted for ------------------------------
    _section_title(document, "Mission")
    mission = document.add_table(rows=0, cols=2)
    mission.style = "Table Grid"
    for label, placeholder in (
        ("Intitulé", "{{ mission_titre }}"),
        ("Acheteur", "{{ mission_acheteur }}"),
        ("Référence", "{{ mission_reference }}"),
        ("Pays", "{{ mission_pays }}"),
        ("Date limite", "{{ mission_echeance }}"),
    ):
        cells = mission.add_row().cells
        cells[0].paragraphs[0].add_run(label).bold = True
        cells[1].text = placeholder

    # --- education ---------------------------------------------------------
    # A loop over a variable the platform declares but does not yet produce.
    # It renders as nothing today and fills itself the day the structured
    # timeline lands, with no change to this file — which is the point of
    # declaring the unavailable variables in the catalogue rather than omitting
    # them.
    _section_title(document, "Formation")
    # `{%p %}` is docxtpl's paragraph-level tag: the paragraph holding the tag
    # is removed and the block between them repeats as whole paragraphs. With
    # an inline `{% for %}` every entry renders into one continuous run, and
    # the first real generation showed one role's missions colliding with the
    # next role's dates.
    document.add_paragraph("{%p for f in formations %}")
    document.add_paragraph("{{ f.annee }} — {{ f.diplome }}, {{ f.etablissement }}")
    document.add_paragraph("{%p endfor %}")

    # --- languages and skills ---------------------------------------------
    _section_title(document, "Langues")
    document.add_paragraph("{{ langues | join(', ') }}")

    _section_title(document, "Compétences techniques")
    document.add_paragraph("{{ competences | join(' · ') }}")

    _section_title(document, "Certifications")
    document.add_paragraph("{{ certifications | join(', ') }}")

    # --- career ------------------------------------------------------------
    _section_title(document, "Expérience professionnelle")
    document.add_paragraph("{%p for e in experiences %}")
    header = document.add_paragraph()
    header.add_run(
        "{{ e.debut }} – {{ e.fin }} · {{ e.poste }}"
        "{% if e.employeur %}, {{ e.employeur }}{% endif %}"
    ).bold = True
    document.add_paragraph("{{ e.missions }}")
    document.add_paragraph("{%p endfor %}")

    # --- provenance --------------------------------------------------------
    # Not decoration. A CV produced by a platform should say what produced it
    # and from which selection, so that a reader who doubts a claim knows where
    # to go and ask.
    document.add_paragraph()
    footer = document.add_paragraph()
    run = footer.add_run(
        "Généré le {{ genere_le }} · sélection rang {{ rang }} · "
        "score de rapprochement {{ score }}"
    )
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    destination.parent.mkdir(parents=True, exist_ok=True)
    # Written to a temporary name and renamed, so an interrupted run never
    # leaves a truncated .docx that opens and renders nonsense.
    staging = destination.with_suffix(".docx.part")
    document.save(staging)
    staging.replace(destination)
    return destination


def _section_title(document, text: str) -> None:
    from docx.shared import Pt

    paragraph = document.add_paragraph()
    run = paragraph.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(11)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", type=Path, default=DEFAULT_OUT)
    args = parser.parse_args()

    path = build(args.out)
    size = path.stat().st_size

    # Prove it is a usable template rather than merely a file that was written:
    # parse it back and report the placeholders it asks for, checked against
    # the catalogue the platform can actually fill.
    from app.services.templates import check_variables, discover_variables

    found = discover_variables(path.read_bytes())
    unknown, unavailable = check_variables(found, kind="cv")

    print(f"écrit      {path} ({size:,} octets)")
    print(f"variables  {', '.join(sorted(found))}")
    if unavailable:
        print(f"en attente {', '.join(unavailable)} — rendus vides jusqu'à l'étape 2")
    if unknown:
        print(f"INCONNUES  {', '.join(unknown)}")
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
