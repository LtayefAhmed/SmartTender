"""Nothing collected should be lost in silence.

Three losses were found by reading the interface rather than by any failure:
attachments arrived as ZIP archives nobody opened, the richest document was
linked from inside the publication and never followed, and a document cap
dropped the specification while keeping the forms. None of them raised.

These tests pin the fixes and, just as importantly, pin the *bounds* — an
archive is attacker-controlled input, and an unbounded reader is a denial of
service waiting to be published as a cahier des charges.
"""

from __future__ import annotations

import io
import zipfile

import pytest

from app.connectors.parsing.links import harvest_document_links
from app.services.extraction import DOCUMENT_MARKER, document_priority, get_extractor


def _zip(members: dict[str, bytes | str]) -> bytes:
    """Build an archive. Deflated, because a stored archive cannot be a bomb."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, payload in members.items():
            archive.writestr(name, payload)
    return buffer.getvalue()


class TestAnArchiveIsOpened:
    """A buyer publishing the whole dossier as one ZIP must not defeat us."""

    def test_every_member_is_read(self):
        content = _zip(
            {
                "dossier/CCTP.txt": "tierce maintenance applicative Java Angular",
                "dossier/ATTRI1.txt": "formulaire administratif",
            }
        )
        result = get_extractor().extract(content, filename="dossier.zip")

        assert result.ok
        assert "tierce maintenance applicative" in result.text
        assert "formulaire administratif" in result.text

    def test_each_member_is_named_in_the_text(self):
        """Provenance survives the archive.

        A 440 000-character tender is far too long to embed as one vector, so
        it will be chunked. A chunk is only worth ranking if you know it came
        from the CCTP and not from a privacy notice.
        """
        result = get_extractor().extract(_zip({"a/CCTP.txt": "prestations"}), filename="d.zip")

        assert f"{DOCUMENT_MARKER.strip()} a/CCTP.txt" in result.text

    def test_the_specification_is_read_before_the_forms(self):
        """Ordering matters only because the member cap can bind."""
        result = get_extractor().extract(
            _zip({"zzz_CCTP.txt": "specification", "aaa_DC1.txt": "formulaire"}),
            filename="dossier.zip",
        )

        assert result.text.index("specification") < result.text.index("formulaire")

    def test_a_docx_is_not_mistaken_for_an_archive(self):
        """A .docx *is* a ZIP. Reading it as one yields a bag of XML parts."""
        from docx import Document

        buffer = io.BytesIO()
        document = Document()
        document.add_paragraph("cahier des clauses techniques particulieres")
        document.save(buffer)

        result = get_extractor().extract(buffer.getvalue(), filename="cctp.docx")

        assert "cahier des clauses techniques" in result.text
        assert "word/document.xml" not in result.text


class TestAnArchiveCannotBeUsedAgainstUs:
    """Bounds on hostile input. Each is independent because each can be
    exhausted on its own."""

    def test_a_nested_archive_is_opened_once(self):
        inner = _zip({"annexe/reglement.txt": "reglement de consultation"})
        result = get_extractor().extract(_zip({"annexes.zip": inner}), filename="dossier.zip")

        assert "reglement de consultation" in result.text

    def test_recursion_stops_at_the_configured_depth(self):
        extractor = get_extractor()
        deepest = _zip({"payload.txt": "trop profond"})
        for level in range(extractor.archive_max_depth + 1):
            deepest = _zip({f"level{level}.zip": deepest})

        result = extractor.extract(deepest, filename="poupees.zip")

        assert "trop profond" not in result.text

    def test_a_bomb_is_stopped_by_the_size_budget(self, monkeypatch):
        """A zip bomb is a few kilobytes that decompress to gigabytes.

        Highly compressible content is what makes it cheap to publish, so the
        defence cannot be the archive's own size — it has to be the running
        total of what has been decompressed.
        """
        extractor = get_extractor()
        monkeypatch.setattr(extractor, "archive_max_total_bytes", 2_000)
        content = _zip({"bomb.txt": "x" * 500_000})

        assert len(content) < 2_000  # a small archive, by design
        result = extractor.extract(content, filename="bomb.zip")

        assert "xxxx" not in result.text
        assert any("budget exhausted" in warning for warning in result.warnings)

    def test_the_budget_keeps_what_it_can_before_stopping(self, monkeypatch):
        """Partial content plus a visible warning beats a silent nothing."""
        extractor = get_extractor()
        monkeypatch.setattr(extractor, "archive_max_total_bytes", 3_000)
        content = _zip({"CCTP.txt": "specification technique", "gros.txt": "y" * 500_000})

        result = extractor.extract(content, filename="melange.zip")

        assert "specification technique" in result.text
        assert any("budget exhausted" in warning for warning in result.warnings)

    def test_exceeding_the_member_cap_is_reported_not_hidden(self, monkeypatch):
        extractor = get_extractor()
        monkeypatch.setattr(extractor, "archive_max_members", 2)
        content = _zip({f"file{index}.txt": f"contenu {index}" for index in range(5)})

        result = extractor.extract(content, filename="beaucoup.zip")

        assert result.ok
        assert any("only the 2 most relevant" in warning for warning in result.warnings)

    def test_ocr_is_budgeted_across_the_whole_archive(self, monkeypatch):
        """The OCR bound was per document, which was enough while a document
        was the unit of work. A dossier of twenty scanned PDFs spent twenty
        times the intended budget and occupied a worker for close to an hour.
        """
        from pypdf import PdfWriter

        def blank_pdf(pages: int) -> bytes:
            writer = PdfWriter()
            for _ in range(pages):
                writer.add_blank_page(width=200, height=200)
            buffer = io.BytesIO()
            writer.write(buffer)
            return buffer.getvalue()

        extractor = get_extractor()
        monkeypatch.setattr(extractor, "max_ocr_pages", 3)
        monkeypatch.setattr(extractor, "ocr_enabled", True)
        monkeypatch.setattr(extractor, "_ensure_ocr", lambda: True)

        # Tesseract itself is not under test; how many pages reach it is.
        requested: list[int] = []

        def spy(content, page_indices):
            requested.append(len(page_indices))
            return dict.fromkeys(page_indices, "texte reconnu")

        monkeypatch.setattr(extractor, "_ocr_pages", spy)

        # Blank pages carry no text layer, so every one of them asks for OCR.
        archive = _zip({"scan_a.pdf": blank_pdf(4), "scan_b.pdf": blank_pdf(4)})
        result = extractor.extract(archive, filename="scans.zip")

        assert sum(requested) == 3, "the allowance is shared, not granted per member"
        assert any("only 3 were processed" in warning for warning in result.warnings)

    def test_a_lone_document_still_gets_the_full_allowance(self, monkeypatch):
        """Sharing the budget must not shrink the ordinary single-PDF case."""
        from pypdf import PdfWriter

        writer = PdfWriter()
        for _ in range(5):
            writer.add_blank_page(width=200, height=200)
        buffer = io.BytesIO()
        writer.write(buffer)

        extractor = get_extractor()
        monkeypatch.setattr(extractor, "max_ocr_pages", 4)
        monkeypatch.setattr(extractor, "ocr_enabled", True)
        monkeypatch.setattr(extractor, "_ensure_ocr", lambda: True)
        requested: list[int] = []
        monkeypatch.setattr(
            extractor,
            "_ocr_pages",
            lambda content, pages: (requested.append(len(pages)), {})[1],
        )

        extractor.extract(buffer.getvalue(), filename="scan.pdf")

        assert requested == [4]

    def test_an_unreadable_archive_fails_without_raising(self):
        result = get_extractor().extract(b"PK\x03\x04 corrompu", filename="casse.zip")

        assert not result.ok
        assert result.error


class TestFormatsTheCompletenessReportExposed:
    """These were not guessed. The completeness endpoint listed them as
    "Unsupported document type" on real consultations already in the corpus —
    which is the whole argument for measuring what is missing."""

    def test_a_7z_dossier_is_read(self):
        """Which archiver the buyer used is not a property of the tender."""
        import py7zr

        buffer = io.BytesIO()
        with py7zr.SevenZipFile(buffer, "w") as archive:
            # py7zr takes the data first, unlike zipfile.
            archive.writestr("recette fonctionnelle du SI", "NP_CCTP_Tester_SI.txt")
            archive.writestr("formulaire", "DC1.txt")

        result = get_extractor().extract(buffer.getvalue(), filename="dossier.7z")

        assert "recette fonctionnelle du SI" in result.text
        assert result.text.index("recette") < result.text.index("formulaire")

    def test_a_spreadsheet_is_read(self):
        """An annexe named "Exigences, pénalités, livrables et indicateurs" is
        the requirements matrix, and it was going unread."""
        from openpyxl import Workbook

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Exigences"
        sheet.append(["Exigence", "Cible"])
        sheet.append(["Disponibilite du service", "99,5 %"])
        buffer = io.BytesIO()
        workbook.save(buffer)

        result = get_extractor().extract(buffer.getvalue(), filename="CCAP_Annexe2.xlsx")

        assert "Exigences" in result.text
        # A requirement and its target must stay on one line, or chunking will
        # separate "disponibilité" from "99,5 %" and neither will mean anything.
        line = next(li for li in result.text.splitlines() if "Disponibilite" in li)
        assert "99,5 %" in line

    def test_a_spreadsheet_yields_values_not_formulas(self):
        """A matcher needs "99,5 %", not "=B4*C4"."""
        from openpyxl import Workbook

        workbook = Workbook()
        workbook.active["A1"] = "=1+1"
        buffer = io.BytesIO()
        workbook.save(buffer)

        result = get_extractor().extract(buffer.getvalue(), filename="calcul.xlsx")

        assert "=1+1" not in result.text


class TestLinksInsideAPublicationAreFollowed:
    """The richest document is often linked from the prose, not listed as an
    attachment. Flattening the page to text destroyed the href."""

    PAGE = """
    <html><body>
      <a href="#haut">Haut de page</a>
      <a href="mailto:acheteur@ministere.fr">Contacter l'acheteur</a>
      <a href="/mentions-legales">Mentions legales</a>
      <a href="/connexion">Se connecter</a>
      <a href="/files/RC_VF.pdf">Reglement de consultation - 1,1 Mo</a>
      <a href="/download/8821">Cahier des clauses techniques</a>
      <a href="/files/dossier.zip">Dossier complet (12,4 Mo)</a>
      <a href="/files/RC_VF.pdf#page=4">Reglement, article 4</a>
    </body></html>
    """

    @pytest.fixture
    def links(self):
        return harvest_document_links(self.PAGE, base_url="https://j360-ext.info/pub/9911")

    def test_documents_are_found(self, links):
        assert {link.url for link in links} == {
            "https://j360-ext.info/files/RC_VF.pdf",
            "https://j360-ext.info/download/8821",
            "https://j360-ext.info/files/dossier.zip",
        }

    def test_navigation_is_not_collected(self, links):
        """Every notice has a login page and a privacy policy. Collecting them
        would cost a download per notice and add nothing."""
        joined = " ".join(link.url for link in links)

        assert "mentions-legales" not in joined
        assert "connexion" not in joined
        assert "mailto" not in joined

    def test_an_opaque_href_is_kept_when_the_label_names_a_document(self, links):
        """`/download/8821` has no extension. Requiring one would miss the way
        most portals serve their files."""
        opaque = next(link for link in links if link.url.endswith("8821"))

        assert opaque.reason == "download-path"
        assert opaque.label == "Cahier des clauses techniques"

    def test_a_fragment_is_not_a_second_document(self, links):
        """`RC_VF.pdf#page=4` is a position inside a file already collected."""
        assert sum(1 for link in links if "RC_VF.pdf" in link.url) == 1

    def test_a_relative_link_without_a_base_is_dropped(self):
        """Recording a URL that cannot be fetched is worse than recording
        nothing: it becomes a permanent failure in the completeness report."""
        assert harvest_document_links('<a href="/files/RC.pdf">RC</a>') == []


class TestImportanceOrdersWhatSurvivesACap:
    """A cap of 10 on a 15-file consultation dropped the règlement de
    consultation, the CCTP and the technical stack — the three that mattered —
    purely by list position."""

    @pytest.mark.parametrize(
        "name",
        ["CNSO_CCTP_VF.pdf", "RC_VF.pdf", "cahier des charges.docx", "DCE_2026.zip"],
    )
    def test_substantive_documents_rank_first(self, name):
        assert document_priority(name) == 0

    @pytest.mark.parametrize("name", ["ATTRI1.pdf", "DC1.doc", "DC2.doc", None])
    def test_administrative_forms_rank_last(self, name):
        assert document_priority(name) == 1
