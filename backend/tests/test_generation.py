"""Producing a document from a validated selection.

The interesting tests are the refusals. A generator that always produces
something is a generator that will one day produce a CV for a person nobody
selected, with an employer nobody wrote, under a heading nobody approved — and
none of those failures announce themselves.
"""

from __future__ import annotations

import io
from datetime import datetime, timezone
from types import SimpleNamespace

import pytest

from app.services.generation import (
    WATERMARK,
    GenerationRefused,
    build_cv_context,
    generate_cv,
)


@pytest.fixture
def template() -> bytes:
    import tempfile
    from pathlib import Path

    from scripts.build_default_template import build

    with tempfile.TemporaryDirectory() as folder:
        return build(Path(folder) / "cv.docx").read_bytes()


def _cv(**overrides):
    base = {
        "original_filename": "cv_dupont.pdf",
        "display_name": "Marie Dupont",
        "headline": "Ingénieure études et développement",
        "criteria": {
            "technologies": ["Java", "Spring"],
            "languages": ["français", "anglais"],
            "certifications": ["AWS Certified"],
            "education_label": "Bac+5 (Master / Ingénieur)",
        },
        "structure": {
            "status": "ok",
            "experiences": [
                {
                    "debut": "2019",
                    "fin": "2022",
                    "poste": "Développeuse Java",
                    "employeur": None,
                    "lieu": None,
                    "missions": "Migration du socle applicatif.",
                }
            ],
            "formations": [
                {
                    "annee": "2018",
                    "diplome": "Diplôme d'ingénieur",
                    "etablissement": "ENSI",
                }
            ],
        },
    }
    base.update(overrides)
    return SimpleNamespace(**base)


def _entry(**overrides):
    base = {
        "label": "Marie Dupont",
        "rank": 1,
        "score": 0.81,
        "decision": "retained",
        "evidence": [{"requirement": 1, "score": 0.7, "passage": "Développement Java 17"}],
    }
    base.update(overrides)
    return SimpleNamespace(**base)


def _tender():
    return SimpleNamespace(
        title="Tierce maintenance applicative",
        buyer="Ministère des Technologies",
        reference="AO 01/2026",
        country="Tunisie",
        deadline=datetime(2026, 3, 31, tzinfo=timezone.utc),
    )


class TestNothingIsGeneratedForSomeoneNobodyChose:
    def test_a_pending_profile_is_refused(self):
        """Left pending means nobody looked. Producing a CV would put a real
        person's name on a real submission on the strength of a rank."""
        with pytest.raises(GenerationRefused):
            generate_cv(
                template_bytes=b"", cv=_cv(), entry=_entry(decision="pending")
            )

    def test_a_rejected_profile_is_refused(self):
        with pytest.raises(GenerationRefused):
            generate_cv(
                template_bytes=b"", cv=_cv(), entry=_entry(decision="rejected")
            )

    def test_a_retained_profile_is_produced(self, template):
        document = generate_cv(
            template_bytes=template, cv=_cv(), entry=_entry(), tender=_tender()
        )

        assert document.content
        assert "Marie Dupont" in _read(document.content)


class TestTheWatermarkIsOnTheFaceOfTheDocument:
    def test_an_unapproved_version_carries_it(self, template):
        document = generate_cv(
            template_bytes=template, cv=_cv(), entry=_entry(), tender=_tender()
        )

        assert WATERMARK in _read(document.content)

    def test_only_an_explicit_approval_removes_it(self, template):
        document = generate_cv(
            template_bytes=template,
            cv=_cv(),
            entry=_entry(),
            tender=_tender(),
            approved=True,
        )

        assert WATERMARK not in _read(document.content)


class TestTheContextIsBuiltFromStoredDataOnly:
    def test_the_timeline_reaches_the_document(self, template):
        text = _read(
            generate_cv(
                template_bytes=template, cv=_cv(), entry=_entry(), tender=_tender()
            ).content
        )

        assert "Développeuse Java" in text
        assert "Migration du socle applicatif." in text
        assert "Diplôme d'ingénieur" in text

    def test_a_missing_employer_stays_missing(self):
        """"We could not read it" and "there wasn't one" are different facts,
        and only the first is true. A plausible default would make the document
        assert something nobody wrote."""
        context, _, _ = build_cv_context(cv=_cv(), entry=_entry())

        assert context["experiences"][0]["employeur"] == ""

    def test_the_label_the_validator_saw_wins_over_a_fresh_lookup(self):
        """A CV re-imported under a new display name must not silently change
        the name on a document produced from an approved selection."""
        context, _, _ = build_cv_context(
            cv=_cv(display_name="M. DUPONT (v2)"), entry=_entry(label="Marie Dupont")
        )

        assert context["nom"] == "Marie Dupont"

    def test_the_evidence_travels_with_the_document(self):
        context, _, _ = build_cv_context(cv=_cv(), entry=_entry())

        assert context["extraits"] == ["Développement Java 17"]

    def test_a_cv_with_no_timeline_still_produces_a_document(self, template):
        """29% of the base yields no degree and 9% no role. A generator that
        refused those would be unusable on our own data — the headings stand
        empty, which is a hole a reviewer can see and fill."""
        document = generate_cv(
            template_bytes=template,
            cv=_cv(structure={}),
            entry=_entry(),
            tender=_tender(),
        )

        assert document.content
        assert "experiences" in document.empty_fields

    def test_the_roles_are_capped(self):
        """A funder's form has a box, not a chapter. The extractor has been
        measured returning twenty roles for one CV."""
        many = {
            "experiences": [
                {"debut": str(year), "fin": str(year + 1), "poste": f"Poste {year}"}
                for year in range(1990, 2020)
            ],
            "formations": [],
        }
        context, _, _ = build_cv_context(cv=_cv(structure=many), entry=_entry())

        assert len(context["experiences"]) == 8


class TestTheFileIsNamedForAHuman:
    def test_the_name_carries_the_person_and_the_tender(self, template):
        document = generate_cv(
            template_bytes=template, cv=_cv(), entry=_entry(), tender=_tender()
        )

        assert document.filename == "Marie_Dupont_AO_01_2026.docx"

    def test_accents_and_punctuation_do_not_reach_the_filesystem(self, template):
        document = generate_cv(
            template_bytes=template,
            cv=_cv(),
            entry=_entry(label="Amél/ie Ben Saïd"),
            tender=_tender(),
        )

        assert document.filename.startswith("Amel_ie_Ben_Said")


def _read(content: bytes) -> str:
    from docx import Document

    document = Document(io.BytesIO(content))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        parts.extend(cell.text for row in table.rows for cell in row.cells)
    return "\n".join(parts)
